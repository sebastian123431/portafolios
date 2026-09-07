
import base64
import csv
import io
import os
import tempfile
import zipfile
from datetime import date, timedelta
from pathlib import Path
from unittest import skipUnless
from unittest.mock import patch

import cv2
import numpy as np
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.files.base import ContentFile
from django.contrib.staticfiles import finders
from django.test import Client, TestCase, override_settings
from django.urls import resolve, reverse
from docx import Document
from openpyxl import Workbook, load_workbook

from .models import ArchivoContrato, Asistencia, Cargo, Contrato, ContratoTrabajador, Login, trabajadores
from .services.documents import DocumentConversionUnavailable, convert_to_pdf
from .urls import urlpatterns


@override_settings(SENDGRID_API_KEY="", ALLOWED_HOSTS=["testserver", "localhost", "127.0.0.1"])
class IntegrationTests(TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        settings_override = override_settings(MEDIA_ROOT=temporary.name)
        settings_override.enable()
        self.addCleanup(settings_override.disable)
        self.media = Path(temporary.name)
        self.worker_data = {
            "nombre": "Persona", "apellido": "Ficticia", "rut": "DEMO-001",
            "direccion": "Calle de prueba", "telefono": "000000000",
            "afp": "AFP demo", "estado_civil": "Soltero", "fecha_nacimiento": "1990-01-02",
            "previcion_salud": "Salud demo", "correo": "persona@example.com",
        }
        self.worker = trabajadores.objects.create(**self.worker_data)
        self.cargo = Cargo.objects.create(
            trabajadores=self.worker, labor="Cargo demo", fecha_ingreso=date.today(), sueldo=1000,
        )
        document = Document()
        document.add_paragraph("Contrato de @nombre@ @apellido@")
        buffer = io.BytesIO()
        document.save(buffer)
        self.docx = buffer.getvalue()
        self.contract = Contrato.objects.create(
            nombre="Contrato ficticio", descripcion="Solo pruebas",
            archivo=SimpleUploadedFile("demo.docx", self.docx),
        )

    def url(self, name, *args):
        return reverse(f"contratos_agiles:{name}", args=args)

    def test_portfolio_and_all_module_routes(self):
        for url in ["/", "/bibliografia/"]:
            with self.subTest(url=url):
                self.assertEqual(self.client.get(url).status_code, 200)
        expected = {
            "guardar_trabajador": 401, "guardar_contrato": 400,
            "generar_contrato": 405, "editar_contrato": 302,
            "cambiar_vigencia": 302, "procesar_datos": 302,
        }
        for pattern in urlpatterns:
            args = []
            for parameter in pattern.pattern.converters:
                args.append("demo.docx" if parameter == "archivo" else self.worker.id)
                if parameter == "contrato_id":
                    args[-1] = self.contract.id
            url = self.url(pattern.name, *args)
            with self.subTest(route=pattern.name):
                self.assertTrue(url.startswith("/proyectos/contratos-agiles/"))
                self.assertEqual(resolve(url).namespace, "contratos_agiles")
                response = self.client.get(url)
                self.assertEqual(response.status_code, expected.get(pattern.name, 200))
                self.assertEqual(response.headers.get("X-Frame-Options"), "SAMEORIGIN")

    def test_registration_login_and_csrf(self):
        client = Client(enforce_csrf_checks=True)
        url = self.url("register")
        self.assertEqual(client.post(url, {}).status_code, 403)
        client.get(url)
        data = {"username": "demo_test", "password": "test-only", "csrfmiddlewaretoken": client.cookies["csrftoken"].value}
        self.assertEqual(client.post(url, data).json()["status"], "success")
        response = client.post(self.url("login"), data, HTTP_X_REQUESTED_WITH="XMLHttpRequest")
        self.assertEqual(response.json()["redirect_url"], self.url("menu"))
        data["password"] = "incorrecta"
        self.assertEqual(client.post(self.url("login"), data, HTTP_X_REQUESTED_WITH="XMLHttpRequest").json()["status"], "error")

    def test_attendance_login_redirect(self):
        data = {"username": "asistencia", "password": "asistencia"}
        self.client.post(self.url("register"), data)
        response = self.client.post(self.url("login"), data, HTTP_X_REQUESTED_WITH="XMLHttpRequest")
        self.assertEqual(response.json()["redirect_url"], self.url("registrar_asistencia"))

    def test_demo_login_is_visible_and_allows_natural_user_access(self):
        response = self.client.get(self.url("login"))

        self.assertContains(response, "Acceso demo")
        self.assertContains(response, "visitante")
        self.assertContains(response, "portafolio2026")
        self.assertTrue(Login.objects.filter(username="visitante", password="portafolio2026").exists())

        login_response = self.client.post(
            self.url("login"),
            {"username": "visitante", "password": "portafolio2026"},
            HTTP_X_REQUESTED_WITH="XMLHttpRequest",
        )
        self.assertEqual(login_response.json()["status"], "success")
        self.assertEqual(login_response.json()["redirect_url"], self.url("menu"))

    def test_search_urls_and_autocomplete(self):
        data = self.client.get(self.url("buscar_trabajadores"), {"q": "Ficticia"}).json()["trabajadores"]
        self.assertEqual(len(data), 1)
        self.assertEqual(data[0]["perfil_url"], self.url("perfil_trabajador", self.worker.id))
        self.assertEqual(data[0]["cargo_url"], self.url("asignar_cargo", self.worker.id))
        self.assertEqual(self.client.get(self.url("autocomplete_trabajadores"), {"term": "Persona"}).json()[0]["id"], self.worker.id)

    def test_worker_registration_and_profile_edit(self):
        data = dict(self.worker_data, nombre="Nueva", prevision_salud="Salud demo")
        response = self.client.post(self.url("guardar_trabajador"), data)
        self.assertRedirects(response, self.url("menu"))
        self.assertEqual(trabajadores.objects.count(), 2)
        data = dict(self.worker_data, nombre="Editada")
        response = self.client.post(self.url("editar_perfil", self.worker.id), data)
        self.assertRedirects(response, self.url("perfil_trabajador", self.worker.id))
        self.worker.refresh_from_db()
        self.assertEqual(self.worker.nombre, "Editada")

    def test_assign_cargo_without_email_configuration(self):
        response = self.client.post(self.url("asignar_cargo", self.worker.id), {
            "labor": "Nuevo cargo", "fecha_ingreso": "2026-01-01", "sueldo": 1000, "vigente": "Vigente",
        }, follow=True)
        self.assertEqual(Cargo.objects.filter(trabajadores=self.worker).count(), 2)
        self.assertContains(response, "Cargo guardado")
        self.assertContains(response, "SendGrid")

    def test_reports_preserve_context(self):
        response = self.client.get(self.url("reportes"))
        self.assertEqual(response.context["previcion_salud_data"], [{"previcion_salud": "Salud demo", "count": 1}])
        self.assertEqual(response.context["afp_data"], [{"afp": "AFP demo", "count": 1}])
        self.assertEqual(response.context["nacimiento_data"], [{"year": "1990", "count": 1}])

    def test_excel_export(self):
        response = self.client.get(self.url("exportar_trabajadores"))
        workbook = load_workbook(io.BytesIO(response.content))
        values = list(workbook.active.values)
        self.assertEqual(values[0], ("id", "nombre", "apellido", "rut", "direccion", "telefono", "afp", "estado_civil", "fecha_nacimiento", "previcion_salud"))
        self.assertEqual(values[1][1:4], ("Persona", "Ficticia", "DEMO-001"))

    def test_csv_import_mapping_and_save(self):
        buffer = io.StringIO()
        writer = csv.DictWriter(buffer, fieldnames=list(self.worker_data))
        writer.writeheader()
        writer.writerow(dict(self.worker_data, nombre="Importada"))
        response = self.client.post(self.url("importar_trabajadores"), {
            "file": SimpleUploadedFile("demo.csv", buffer.getvalue().encode())
        })
        self.assertContains(response, "Mapeo de columnas")
        mapping = {key: key for key in self.worker_data}
        response = self.client.post(self.url("importar_trabajadores"), {
            **mapping, "guardar_mapeo": "1", "csrfmiddlewaretoken": "test",
        }, follow=True)
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Importada")
        response = self.client.post(self.url("importar_trabajadores"), {"guardar": "1"})
        self.assertRedirects(response, self.url("visualizar_trabajadores"))
        self.assertTrue(trabajadores.objects.filter(nombre="Importada").exists())

    def test_excel_import_serializes_dates_in_session(self):
        workbook = Workbook()
        workbook.active.append(list(self.worker_data))
        row = dict(self.worker_data, fecha_nacimiento=date(1990, 1, 2))
        workbook.active.append(list(row.values()))
        buffer = io.BytesIO()
        workbook.save(buffer)
        response = self.client.post(self.url("importar_trabajadores"), {"file": SimpleUploadedFile("demo.xlsx", buffer.getvalue())})
        self.assertEqual(response.status_code, 200)
        self.assertIsInstance(self.client.session["df"]["fecha_nacimiento"][0], str)

    def test_contract_upload_preview_and_save_docx(self):
        response = self.client.post(self.url("subir_contrato"), {
            "nombre": "Nueva plantilla", "descripcion": "Ficticia", "archivo": SimpleUploadedFile("otra.docx", self.docx),
        })
        self.assertRedirects(response, self.url("listar_contratos"))
        response = self.client.post(self.url("previsualizar_contrato"), {"contrato": SimpleUploadedFile("preview.docx", self.docx)})
        self.assertContains(response, "Contrato de @nombre@")
        response = self.client.get(self.url("guardar_contrato"), {"file_path": response.context["file_path"]})
        self.assertEqual(response.status_code, 200)
        self.assertTrue((self.media / "contratos_finales/contrato_generado.docx").exists())

    def test_document_service_rejects_invalid_docx_and_external_paths(self):
        response = self.client.post(self.url("previsualizar_contrato"), {
            "contrato": SimpleUploadedFile("invalid.docx", b"not a docx"),
        })
        self.assertContains(response, "DOCX válido", status_code=400)
        response = self.client.get(self.url("guardar_contrato"), {"file_path": "../outside.docx"})
        self.assertContains(response, "ruta del documento", status_code=400)

    def test_preview_images_and_html_text(self):
        from PIL import Image
        from docx.shared import Inches
        png = io.BytesIO()
        Image.new("RGB", (24, 24), "blue").save(png, format="PNG")
        document = Document()
        document.add_paragraph("Antes <script>texto</script>")
        document.add_picture(io.BytesIO(png.getvalue()), width=Inches(1))
        document.add_paragraph("Después")
        buffer = io.BytesIO()
        document.save(buffer)
        response = self.client.post(self.url("previsualizar_contrato"), {
            "contrato": SimpleUploadedFile("image.docx", buffer.getvalue()),
        })
        self.assertContains(response, "&lt;script&gt;texto&lt;/script&gt;")
        html = response.context["html_content"]
        self.assertLess(html.index("Antes"), html.index("<img"))
        self.assertLess(html.index("<img"), html.index("Después"))
        self.assertEqual(len(list((self.media / "tmp").rglob("*.png"))), 1)

    def test_saved_preview_replaces_markers_and_inserts_logo(self):
        from PIL import Image
        Image.new("RGB", (24, 24), "blue").save(self.media / "logo_empresa.png")
        document = Document()
        document.add_paragraph("@nombre@ — @fecha@")
        document.add_paragraph("@logo@")
        source = self.media / "preview.docx"
        document.save(source)
        response = self.client.get(self.url("guardar_contrato"), {"file_path": "preview.docx"})
        self.assertEqual(response.status_code, 200)
        saved = Document(self.media / "contratos_finales/contrato_generado.docx")
        self.assertEqual(saved.paragraphs[0].text, "Juan Pérez — 01/11/2024")
        self.assertEqual(len(saved.inline_shapes), 1)

    @skipUnless(os.getenv("CONTRATOS_TEST_REAL_WORD") == "1", "Prueba optativa con Microsoft Word instalado")
    def test_real_word_individual_specific_and_bulk_generation(self):
        for name, args in [("generar_contrato", []), ("generar_contrato_especifico", [self.worker.id]), ("generar_contratos_masivos", [])]:
            with self.subTest(route=name):
                response = self.client.post(self.url(name, *args), self.generation_data())
                self.assertEqual(response.status_code, 200, response.content[:300])
                if name == "generar_contratos_masivos":
                    with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
                        self.assertEqual(len(archive.namelist()), 1)
                        pdf = archive.read(archive.namelist()[0])
                else:
                    pdf = response.content
                self.assertTrue(pdf.startswith(b"%PDF-"))
                self.assertIn(b"%%EOF", pdf[-1024:])
                self.assertGreater(len(pdf), 1000)
                generated = Document(self.media / "contratos_finales/contrato_Persona_Ficticia.docx")
                self.assertIn("Persona Ficticia", generated.paragraphs[0].text)
        self.assertTrue(ArchivoContrato.objects.exists())

    def generation_data(self):
        return {
            "trabajador": self.worker.id, "trabajadores": [self.worker.id],
            "contrato": self.contract.id, "fecha_inicio": date.today().isoformat(),
            "fecha_termino": (date.today() + timedelta(days=30)).isoformat(),
        }

    def test_pdf_generation_without_windows_is_controlled(self):
        with patch("proyectos.contratos_agiles.services.documents.sys.platform", "linux"):
            for name, args in [("generar_contrato", []), ("generar_contrato_especifico", [self.worker.id]), ("generar_contratos_masivos", [])]:
                with self.subTest(route=name):
                    response = self.client.post(self.url(name, *args), self.generation_data())
                    self.assertContains(response, "requiere Windows", status_code=503)
            self.assertEqual(self.client.get(self.url("menu")).status_code, 200)
        self.assertEqual(ContratoTrabajador.objects.count(), 0)

    def test_pdf_generation_with_conversion_service(self):
        def fake_pdf(source, destination):
            self.assertIn("Persona Ficticia", Document(source).paragraphs[0].text)
            Path(destination).write_bytes(b"%PDF-1.4\n% fictitious test document")
        with patch("proyectos.contratos_agiles.services.documents.convert_to_pdf", side_effect=fake_pdf):
            response = self.client.post(self.url("generar_contrato_especifico", self.worker.id), self.generation_data())
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertEqual(ContratoTrabajador.objects.count(), 1)
        self.assertEqual(ArchivoContrato.objects.count(), 1)

    def test_email_not_marked_sent_on_failure(self):
        response = self.client.post(self.url("enviar_correos_masivos"))
        self.assertContains(response, "SendGrid", status_code=503)
        self.worker.refresh_from_db()
        self.assertEqual(self.worker.sent, "no_enviado")

    def test_bulk_generation_does_not_duplicate_contracts(self):
        def fake_pdf(source, destination):
            Path(destination).write_bytes(b"%PDF-1.4\n% demo")
        with patch("proyectos.contratos_agiles.services.documents.convert_to_pdf", side_effect=fake_pdf):
            response = self.client.post(self.url("generar_contratos_masivos"), self.generation_data())
        self.assertEqual(response["Content-Type"], "application/zip")
        with zipfile.ZipFile(io.BytesIO(response.content)) as archive:
            self.assertEqual(len(archive.namelist()), 1)
        self.assertEqual(ContratoTrabajador.objects.count(), 1)
        self.assertEqual(ArchivoContrato.objects.count(), 1)

    def test_bulk_generation_handles_existing_individual_duplicates(self):
        for _ in range(2):
            ContratoTrabajador.objects.create(
                trabajador=self.worker, ArchivoContrato=self.contract,
                fecha_termino=date.today() + timedelta(days=30),
            )
        def fake_pdf(source, destination):
            Path(destination).write_bytes(b"%PDF-1.4\n% demo")
        with patch("proyectos.contratos_agiles.services.documents.convert_to_pdf", side_effect=fake_pdf):
            response = self.client.post(self.url("generar_contratos_masivos"), self.generation_data())
        self.assertEqual(response.status_code, 200)
        self.assertEqual(ContratoTrabajador.objects.count(), 2)
        self.assertEqual(ArchivoContrato.objects.count(), 1)

    def test_attendance_with_synthetic_image_and_stub_recognition(self):
        _, encoded = cv2.imencode(".png", np.zeros((160, 160, 3), dtype=np.uint8))
        payload = "data:image/png;base64," + base64.b64encode(encoded).decode()
        with patch("proyectos.contratos_agiles.views.identificar_trabajador", return_value=(self.worker.id, None)):
            for _ in range(2):
                response = self.client.post(self.url("registrar_asistencia_por_imagen"), {"image": payload})
                self.assertEqual(response.json()["status"], "success")
        self.assertEqual(Asistencia.objects.filter(trabajador=self.worker, presente=True).count(), 1)

    def test_real_face_detector_recognizer_and_attendance(self):
        from .views import detectar_rostros
        fixture = Path(__file__).parent / "testdata/synthetic_face.png"
        self.worker.foto.save("synthetic_face.png", ContentFile(fixture.read_bytes()))
        reference = cv2.imread(str(fixture))
        self.assertEqual(len(detectar_rostros(cv2.cvtColor(reference, cv2.COLOR_BGR2GRAY))), 1)
        # Cambiar iluminación y comprimir la captura, sin sustituir detector ni LBPH.
        capture = cv2.convertScaleAbs(reference, alpha=0.92, beta=7)
        _, encoded = cv2.imencode(".jpg", capture, [cv2.IMWRITE_JPEG_QUALITY, 85])
        payload = "data:image/jpeg;base64," + base64.b64encode(encoded).decode()
        for _ in range(2):
            response = self.client.post(self.url("registrar_asistencia_por_imagen"), {"image": payload})
            self.assertEqual(response.json()["status"], "success", response.json())
        self.assertEqual(Asistencia.objects.filter(trabajador=self.worker, presente=True).count(), 1)
        self.assertContains(self.client.get(self.url("ver_asistencia")), "Persona Ficticia")

    def test_real_face_pipeline_rejects_capture_without_face(self):
        _, encoded = cv2.imencode(".png", np.zeros((160, 160, 3), dtype=np.uint8))
        payload = "data:image/png;base64," + base64.b64encode(encoded).decode()
        response = self.client.post(self.url("registrar_asistencia_por_imagen"), {"image": payload})
        self.assertEqual(response.json()["status"], "error")
        self.assertEqual(Asistencia.objects.count(), 0)

    def test_image_upload_uses_storage_url_and_requires_csrf(self):
        client = Client(enforce_csrf_checks=True)
        url = self.url("upload_image")
        self.assertEqual(client.post(url, {}).status_code, 403)
        client.get(self.url("login"))
        response = client.post(url, {
            "csrfmiddlewaretoken": client.cookies["csrftoken"].value,
            "upload": SimpleUploadedFile("demo.png", b"synthetic-image"),
        })
        self.assertTrue(response.json()["url"].startswith("/media/contratos_agiles/images/"))

    def test_static_and_opencv_resources(self):
        for asset in ["css/app.css", "images/logocampana.png", "pdfjs/web/viewer.html"]:
            self.assertIsNotNone(finders.find(f"contratos_agiles/{asset}"))
        from .views import FACE_CASCADE, PROFILE_FACE_CASCADE
        self.assertFalse(FACE_CASCADE.empty())
        self.assertFalse(PROFILE_FACE_CASCADE.empty())
