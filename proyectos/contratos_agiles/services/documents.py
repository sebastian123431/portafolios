"""Conversión de documentos sin cargar COM durante el inicio de Django."""
import sys
import subprocess
from dataclasses import dataclass
from pathlib import Path
from uuid import uuid4

from django.conf import settings
from django.core.files.base import ContentFile
from django.core.files.storage import default_storage
from django.utils.html import escape
from django.utils.text import get_valid_filename
from docx import Document
from docx.shared import Inches


class DocumentProcessingError(ValueError):
    """El documento de entrada no puede procesarse."""


@dataclass(frozen=True)
class GeneratedContract:
    docx_path: Path
    pdf_path: Path


def read_document(source):
    try:
        return Document(source)
    except Exception as exc:
        raise DocumentProcessingError("No se pudo abrir el documento. Selecciona un archivo DOCX válido.") from exc


def preview_uploaded_document(upload):
    """Guarda y previsualiza el DOCX conservando el orden de texto e imágenes."""
    file_path = default_storage.save(f"tmp/{upload.name}", upload)
    document = read_document(default_storage.path(file_path))
    fragments = []
    image_count = 0
    preview_id = uuid4().hex
    for paragraph in document.paragraphs:
        content = ["<p>"]
        for run in paragraph.runs:
            if run.text:
                content.append(f"{escape(run.text)} ")
            for blip in run.element.xpath(".//a:blip"):
                image_id = blip.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                part = document.part.related_parts[image_id]
                image_name = default_storage.save(
                    f"tmp/{preview_id}/image_{image_count}.png", ContentFile(part.blob)
                )
                width, height = 600, 400
                inline_shapes = run._element.xpath(".//a:blip")
                if inline_shapes:
                    ext_elements = inline_shapes[0].getparent().xpath(".//a:ext")
                    if ext_elements:
                        cx, cy = ext_elements[0].get("cx"), ext_elements[0].get("cy")
                        if cx and cy:
                            width, height = int(int(cx) * 0.0001), int(int(cy) * 0.0001)
                content.append(
                    f'<img src="{escape(default_storage.url(image_name))}" alt="Image {image_count}" '
                    f'style="width:{width}px; height:{height}px; vertical-align:middle;">'
                )
                image_count += 1
        content.append("</p>")
        fragments.append("".join(content))
    return {"html_content": "".join(fragments), "file_path": file_path}


def save_preview_document(file_path):
    """Aplica los marcadores de la previsualización y guarda su DOCX final."""
    media_root = Path(settings.MEDIA_ROOT).resolve()
    source = (media_root / file_path).resolve()
    if not source.is_relative_to(media_root):
        raise DocumentProcessingError("La ruta del documento no es válida.")
    if not source.is_file():
        raise FileNotFoundError("El archivo no existe.")
    document = read_document(source)
    replace_placeholders(document, {"nombre": "Juan Pérez", "fecha": "01/11/2024"})
    for paragraph in document.paragraphs:
        if "@logo@" in paragraph.text:
            logo = media_root / "logo_empresa.png"
            if not logo.is_file():
                raise DocumentProcessingError("Falta logo_empresa.png para completar este documento.")
            paragraph.text = paragraph.text.replace("@logo@", "")
            paragraph.add_run().add_picture(str(logo), width=Inches(1.5))
    destination = media_root / "contratos_finales/contrato_generado.docx"
    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(destination)
    return destination


def generate_contract_files(source, context, filename):
    """Lee la plantilla, completa marcadores, guarda DOCX y convierte a PDF."""
    document = read_document(source)
    replace_placeholders(document, context)
    output_dir = Path(settings.MEDIA_ROOT) / "contratos_finales"
    output_dir.mkdir(parents=True, exist_ok=True)
    docx_path = output_dir / get_valid_filename(filename)
    pdf_path = docx_path.with_suffix(".pdf")
    document.save(docx_path)
    convert_to_pdf(docx_path, pdf_path)
    return GeneratedContract(docx_path, pdf_path)


def replace_placeholders(document, context):
    """Conserva la sustitución de marcadores del generador original."""
    for paragraph in document.paragraphs:
        for key, value in context.items():
            if f"@{key}@" in paragraph.text:
                paragraph.text = paragraph.text.replace(f"@{key}@", str(value))
    return document


class DocumentConversionUnavailable(RuntimeError):
    pass


def convert_to_pdf(source, destination):
    """Aísla COM del servidor: un fallo nativo de Word no termina Django."""
    if sys.platform != "win32":
        raise DocumentConversionUnavailable(
            "La conversión DOCX a PDF requiere Windows y Microsoft Word. "
            "El resto de Contratos Ágiles continúa disponible."
        )
    worker = Path(__file__).with_name("word_conversion.py")
    destination = Path(destination).resolve()
    try:
        result = subprocess.run(
            [sys.executable, str(worker), str(Path(source).resolve()), str(destination)],
            capture_output=True, timeout=60, creationflags=subprocess.CREATE_NO_WINDOW,
        )
    except subprocess.TimeoutExpired as exc:
        raise DocumentConversionUnavailable("Microsoft Word excedió el tiempo de conversión. Inténtalo nuevamente.") from exc
    except OSError as exc:
        raise DocumentConversionUnavailable("No se pudo iniciar el proceso de conversión de Microsoft Word.") from exc
    if result.returncode != 0 or not destination.is_file():
        raise DocumentConversionUnavailable(
            "No fue posible convertir el documento a PDF. Comprueba que pywin32 y Microsoft "
            "Word estén instalados y disponibles en esta sesión de Windows."
        )
    return destination
